"""Ngân sách hoạt động 2026: đọc file kế hoạch, nhập chi thực tế, cảnh báo vượt và phân quyền.

Chạy từ thư mục gốc dự án: python -m unittest tests.test_opbudget
"""
import io
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

import openpyxl

from backend import db, opbudget


def plan_workbook(path: Path):
    """File giả cùng bố cục sheet TỔNG HỢP: dòng tiêu đề, mục lớn, nhóm con, hạng mục có mã."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TỔNG HỢP"
    ws.append(["ĐỀ XUẤT NGÂN SÁCH VĂN PHÒNG NĂM 2026"])
    ws.append(["MÃ CÔNG VIỆC", "TÊN CÔNG VIỆC"])

    def put(code="", name="", basic=None, extra=None, total=None, invest=None, note=None):
        row = [None] * 28
        row[0], row[1], row[11], row[15], row[18], row[25], row[16] = code, name, basic, extra, total, invest, note
        ws.append(row)
    put(1.0, "Ngân sách văn phòng", total=999)               # dòng tổng toàn công ty: bỏ qua
    put(2.0, "TỔNG CỘNG CHI PHÍ HÀNH CHÍNH VĂN PHÒNG", total=300)
    put("", "CHI PHÍ VẬN HÀNH ĐỘI XE", total=300)
    put("9300.03.01", "AZB - Xăng dầu", 100, 20, 120_000_000, note="Giảm 20% nếu DT rớt")
    put("9300.03.02", "AZB - Bảo dưỡng", 60, 0, 60_000_000)
    put(3.0, "CHI PHÍ TUYỂN DỤNG", total=24)
    put("9200.01", "AZB - Đăng tuyển", 24, 0, 24_000_000, invest=5)
    put("9200.02", "AZB - Head hunter", 0, 0, 0)
    wb.save(path)


class OpBudgetBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        patch = mock.patch.object(db, "DB_PATH", Path(self.tmp.name) / "azb.db")
        patch.start()
        self.addCleanup(patch.stop)
        self.addCleanup(self.tmp.cleanup)
        db.init_db()
        self.plan = Path(self.tmp.name) / "plan.xlsx"
        plan_workbook(self.plan)


class PlanImportTest(OpBudgetBase):
    def test_parse_sections_groups_and_lines(self):
        lines = opbudget.parse_plan(str(self.plan))
        self.assertEqual([l["code"] for l in lines], ["9300.03.01", "9300.03.02", "9200.01", "9200.02"])
        fuel = lines[0]
        self.assertEqual(fuel["section"], "Chi phí hành chính văn phòng")  # bỏ chữ "TỔNG CỘNG", không viết hoa toàn bộ
        self.assertEqual(fuel["grp"], "Chi phí vận hành đội xe")
        self.assertEqual(fuel["plan_total"], 120_000_000)
        self.assertEqual(fuel["note"], "Giảm 20% nếu DT rớt")
        self.assertEqual(lines[2]["grp"], "")  # mục lớn không có nhóm con
        self.assertEqual(lines[2]["invest"], 5)

    def test_rejects_other_files(self):
        wb = openpyxl.Workbook()
        wb.active.append(["a", "b"])
        other = Path(self.tmp.name) / "other.xlsx"
        wb.save(other)
        with self.assertRaises(ValueError):
            opbudget.parse_plan(str(other))

    def test_reimport_keeps_actuals_and_drops_unused_lines(self):
        with db.connect() as c:
            opbudget.import_plan(c, str(self.plan))
            opbudget.save_actuals(c, None, "9300.03.01", {"1": 5})
            wb = openpyxl.load_workbook(self.plan)
            ws = wb.active
            ws.delete_rows(9, 2)  # file mới bỏ 9200.01 và 9200.02
            ws.delete_rows(6, 1)  # và bỏ 9300.03.01, nhưng hạng mục này đã có số chi nên phải giữ lại
            wb.save(self.plan)
            res = opbudget.import_plan(c, str(self.plan))
            codes = {r[0] for r in c.execute("SELECT code FROM budget_lines")}
        self.assertIn("9300.03.01", codes)       # đã có chi thực tế: không xoá
        self.assertNotIn("9200.01", codes)
        self.assertEqual(res["removed"], 2)


class AlertTest(OpBudgetBase):
    def setUp(self):
        super().setUp()
        with db.connect() as c:
            opbudget.import_plan(c, str(self.plan))

    def status(self, code, months, as_of):
        with db.connect() as c:
            opbudget.save_actuals(c, None, code, months)
            d = opbudget.dashboard(c, as_of)
        return next(l for s in d["sections"] for g in s["groups"] for l in g["lines"] if l["code"] == code)

    def test_thresholds(self):
        # 9300.03.01 có ngân sách 120 triệu
        self.assertEqual(self.status("9300.03.01", {"1": 5_000_000}, 12)["status"], "ok")
        near = self.status("9300.03.01", {"2": 91_000_000}, 12)    # 96 triệu = 80%
        self.assertEqual(near["status"], "near")
        self.assertEqual(near["percent"], 80.0)
        over = self.status("9300.03.01", {"3": 30_000_000}, 12)    # 126 triệu
        self.assertEqual(over["status"], "over")
        self.assertEqual(over["over_by"], 6_000_000)

    def test_exactly_at_plan_is_not_over(self):
        line = self.status("9300.03.01", {"1": 120_000_000}, 12)
        self.assertEqual(line["status"], "near")

    def test_pace_warning_when_spending_faster_than_year_progress(self):
        # đến tháng 3 theo tiến độ được 30 triệu; chi 40 triệu (33% ngân sách, chưa tới 80%)
        self.assertEqual(self.status("9300.03.01", {"1": 40_000_000}, 3)["status"], "pace")
        self.assertEqual(self.status("9300.03.01", {"1": 30_000_000}, 3)["status"], "ok")

    def test_spending_without_budget_is_over(self):
        line = self.status("9200.02", {"1": 1_000_000}, 12)  # ngân sách 0
        self.assertEqual(line["status"], "over")
        self.assertEqual(line["over_by"], 1_000_000)

    def test_months_after_cutoff_are_not_counted(self):
        line = self.status("9300.03.01", {"1": 10_000_000, "9": 500_000_000}, 6)
        self.assertEqual(line["spent"], 10_000_000)
        self.assertEqual(line["year_spent"], 510_000_000)

    def test_dashboard_totals_alerts_and_default_month(self):
        with db.connect() as c:
            opbudget.save_actuals(c, None, "9300.03.01", {"1": 130_000_000})
            opbudget.save_actuals(c, None, "9300.03.02", {"4": 50_000_000})
            d = opbudget.dashboard(c, today=date(2026, 10, 26))
        self.assertEqual(d["as_of"], 4)  # mặc định: tháng gần nhất đã có số chi
        self.assertEqual(d["totals"]["plan"], 204_000_000)
        self.assertEqual(d["totals"]["spent"], 180_000_000)
        self.assertEqual([a["code"] for a in d["alerts"]][0], "9300.03.01")  # vượt xếp trước sắp vượt
        self.assertEqual(d["totals"]["over"], 1)
        self.assertEqual(d["monthly"][0]["spent"], 130_000_000)
        self.assertEqual(d["monthly"][3]["cum_spent"], 180_000_000)

    def test_default_month_without_actuals_is_current_month(self):
        with db.connect() as c:
            self.assertEqual(opbudget.dashboard(c, today=date(2026, 10, 26))["as_of"], 10)

    def test_empty_database_has_no_plan(self):
        with db.connect() as c:
            c.execute("DELETE FROM budget_lines")
            self.assertFalse(opbudget.dashboard(c)["has_plan"])


class ActualsInputTest(OpBudgetBase):
    def setUp(self):
        super().setUp()
        with db.connect() as c:
            opbudget.import_plan(c, str(self.plan))

    def test_blank_clears_month_and_rejects_bad_input(self):
        with db.connect() as c:
            opbudget.save_actuals(c, None, "9200.01", {"1": 10, "2": 20})
            opbudget.save_actuals(c, None, "9200.01", {"1": None, "2": 25})
            got = dict(c.execute("SELECT month, amount FROM budget_actuals WHERE code='9200.01'").fetchall())
            self.assertEqual(got, {2: 25})
            with self.assertRaises(ValueError):
                opbudget.save_actuals(c, None, "9200.01", {"13": 1})
            with self.assertRaises(ValueError):
                opbudget.save_actuals(c, None, "khong-co", {"1": 1})
            with self.assertRaises(ValueError):
                opbudget.save_actuals(c, None, "9200.01", {"1": float("inf")})

    def test_import_wide_csv_skips_blank_cells_and_reports_unknown_codes(self):
        csv = "Mã;Hạng mục;T1;T2;Tháng 3\n9200.01;x;1000;;3000\n9999.99;y;5;5;5\n".encode("utf-8")
        with db.connect() as c:
            opbudget.save_actuals(c, None, "9200.01", {"2": 777})
            res = opbudget.import_actuals(c, None, csv, "a.csv")
            got = dict(c.execute("SELECT month, amount FROM budget_actuals WHERE code='9200.01'").fetchall())
        self.assertEqual(got, {1: 1000, 2: 777, 3: 3000})  # ô trống không xoá số đã có
        self.assertEqual(res["unknown"], ["9999.99"])
        self.assertEqual(res["entries"], 2)

    def test_import_long_format(self):
        csv = "Mã,Tháng,Số tiền\n9200.01,5,1200000\n9200.01,6,300000\n".encode("utf-8")
        with db.connect() as c:
            res = opbudget.import_actuals(c, None, csv, "a.csv")
            total = c.execute("SELECT SUM(amount) FROM budget_actuals").fetchone()[0]
        self.assertEqual((res["entries"], total), (2, 1_500_000))

    def test_template_round_trips_through_import(self):
        with db.connect() as c:
            opbudget.save_actuals(c, None, "9300.03.01", {"1": 111, "12": 222})
            data = opbudget.template_xlsx(c)
            c.execute("DELETE FROM budget_actuals")
            res = opbudget.import_actuals(c, None, data, "mau.xlsx")
            got = dict(c.execute("SELECT month, amount FROM budget_actuals").fetchall())
        self.assertEqual(got, {1: 111, 12: 222})
        self.assertEqual(res["unknown"], [])

    def test_import_without_code_column_is_rejected(self):
        with db.connect() as c, self.assertRaises(ValueError):
            opbudget.import_actuals(c, None, b"a,b\n1,2\n", "a.csv")


class DecisionTest(OpBudgetBase):
    """TGĐ duyệt bổ sung ngân sách / chặn / cho tiếp tục phát sinh."""
    CODE = "9300.03.01"  # ngân sách 120 triệu

    def setUp(self):
        super().setUp()
        with db.connect() as c:
            opbudget.import_plan(c, str(self.plan))

    def line(self, c, as_of=12):
        d = opbudget.dashboard(c, as_of)
        return next(l for s in d["sections"] for g in s["groups"] for l in g["lines"] if l["code"] == self.CODE), d

    def test_supplement_raises_budget_and_clears_overrun(self):
        with db.connect() as c:
            opbudget.save_actuals(c, None, self.CODE, {"1": 150_000_000})
            self.assertEqual(self.line(c)[0]["status"], "over")
            res = opbudget.decide(c, None, self.CODE, "supplement", 40_000_000, "Tăng giá xăng")
            line, d = self.line(c)
        self.assertEqual(res["supplement"], 40_000_000)
        self.assertEqual((line["plan"], line["base_plan"], line["supplement"]), (160_000_000, 120_000_000, 40_000_000))
        self.assertEqual(line["status"], "near")  # 150/160 = 93,75%
        self.assertEqual(d["totals"]["supplement"], 40_000_000)
        self.assertEqual(d["totals"]["plan"], 244_000_000)
        self.assertEqual(d["decisions"][0]["kind"], "supplement")

    def test_supplement_needs_amount_and_reason(self):
        with db.connect() as c:
            for amount, note in ((0, "x"), (-5, "x"), (None, "x"), (1000, ""), (1000, "   ")):
                with self.assertRaises(ValueError, msg=(amount, note)):
                    opbudget.decide(c, None, self.CODE, "supplement", amount, note)
            with self.assertRaises(ValueError):
                opbudget.decide(c, None, self.CODE, "xoa", 1, "x")
            with self.assertRaises(ValueError):
                opbudget.decide(c, None, "khong-co", "block")

    def test_block_rejects_more_spending_but_allows_corrections(self):
        with db.connect() as c:
            opbudget.save_actuals(c, None, self.CODE, {"1": 100})
            opbudget.decide(c, None, self.CODE, "block", note="Dừng chi xăng")
            self.assertTrue(self.line(c)[0]["blocked"])
            with self.assertRaises(ValueError):
                opbudget.save_actuals(c, None, self.CODE, {"2": 50})            # thêm tháng mới
            with self.assertRaises(ValueError):
                opbudget.save_actuals(c, None, self.CODE, {"1": 101})           # sửa tăng
            got = dict(c.execute("SELECT month, amount FROM budget_actuals WHERE code=?", (self.CODE,)).fetchall())
            self.assertEqual(got, {1: 100})                                     # không ghi dở dang
            opbudget.save_actuals(c, None, self.CODE, {"1": 60})                # sửa giảm vẫn được
            self.assertEqual(c.execute("SELECT amount FROM budget_actuals WHERE code=?", (self.CODE,)).fetchone()[0], 60)
            with self.assertRaises(ValueError):
                opbudget.decide(c, None, self.CODE, "block")                    # đã chặn rồi

    def test_unblock_allows_spending_again(self):
        with db.connect() as c:
            opbudget.decide(c, None, self.CODE, "block")
            opbudget.decide(c, None, self.CODE, "unblock", note="Cho tiếp tục")
            opbudget.save_actuals(c, None, self.CODE, {"1": 5})
            self.assertFalse(self.line(c)[0]["blocked"])
            with self.assertRaises(ValueError):
                opbudget.decide(c, None, self.CODE, "unblock")

    def test_supplement_lifts_block(self):
        with db.connect() as c:
            opbudget.decide(c, None, self.CODE, "block")
            self.assertTrue(opbudget.decide(c, None, self.CODE, "supplement", 10, "Bổ sung")["blocked"] is False)

    def test_blocked_line_stays_in_alert_list_and_import_skips_it(self):
        csv = chr(10).join(["Mã;T1", self.CODE + ";5000", "9200.01;7000", ""]).encode("utf-8")
        with db.connect() as c:
            opbudget.decide(c, None, self.CODE, "block")
            d = opbudget.dashboard(c)
            self.assertEqual([a["code"] for a in d["alerts"]], [self.CODE])
            res = opbudget.import_actuals(c, None, csv, "a.csv")
            codes = {r[0] for r in c.execute("SELECT code FROM budget_actuals")}
        self.assertEqual(res["blocked"], [self.CODE])
        self.assertEqual((res["lines"], codes), (1, {"9200.01"}))
        self.assertEqual(d["totals"]["blocked"], 1)


class BudgetEndpointTest(OpBudgetBase):
    def client_for(self, role):
        from fastapi.testclient import TestClient
        from backend import app as app_module, auth
        with db.connect() as c:
            c.execute("INSERT INTO users (id, username, role, password_hash) VALUES (1, 'u', ?, 'x')", (role,))
            token = auth.create_session(c, 1)
        client = TestClient(app_module.app)
        client.cookies.set(auth.SESSION_COOKIE, token)
        return client

    def test_only_admin_and_ceo_can_use_budget(self):
        for role in ("manager", "viewer"):
            db.init_db()
            with db.connect() as c:
                c.execute("DELETE FROM users")
            client = self.client_for(role)
            self.assertEqual(client.get("/api/budget").status_code, 403, role)
            self.assertEqual(client.put("/api/budget/lines/9200.01/actuals", json={"months": {"1": 1}}).status_code, 403, role)
            self.assertEqual(client.get("/api/budget/actuals/template").status_code, 403, role)
            self.assertEqual(client.post("/api/budget/plan/import", files={"file": ("p.xlsx", b"x")}).status_code, 403, role)

    def test_ceo_flow_import_plan_enter_actuals_see_alert(self):
        client = self.client_for("ceo")
        r = client.post("/api/budget/plan/import", files={"file": ("p.xlsx", self.plan.read_bytes())})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["lines"], 4)
        r = client.put("/api/budget/lines/9300.03.01/actuals", json={"months": {"1": 150_000_000}})
        self.assertEqual(r.status_code, 200, r.text)
        d = client.get("/api/budget").json()
        self.assertEqual(d["alerts"][0]["code"], "9300.03.01")
        self.assertEqual(d["alerts"][0]["status"], "over")
        self.assertEqual(client.put("/api/budget/lines/nope/actuals", json={"months": {"1": 1}}).status_code, 400)
        self.assertEqual(client.post("/api/budget/plan/import", files={"file": ("p.txt", b"x")}).status_code, 400)
        with db.connect() as c:
            logged = c.execute("SELECT COUNT(*) FROM audit_log WHERE action IN ('budget_actuals','budget_plan_import')").fetchone()[0]
        self.assertEqual(logged, 2)

    def test_decision_endpoint(self):
        client = self.client_for("ceo")
        client.post("/api/budget/plan/import", files={"file": ("p.xlsx", self.plan.read_bytes())})
        url = "/api/budget/lines/9300.03.01/decision"
        self.assertEqual(client.post(url, json={"kind": "supplement", "amount": 5, "note": ""}).status_code, 400)
        r = client.post(url, json={"kind": "block", "note": "Dừng"})
        self.assertEqual((r.status_code, r.json()["blocked"]), (200, True))
        r = client.put("/api/budget/lines/9300.03.01/actuals", json={"months": {"1": 10}})
        self.assertEqual(r.status_code, 400)
        self.assertIn("chặn", r.json()["detail"])
        r = client.post(url, json={"kind": "supplement", "amount": 30_000_000, "note": "Bổ sung quý 4"})
        self.assertEqual((r.status_code, r.json()["blocked"]), (200, False))
        self.assertEqual(client.put("/api/budget/lines/9300.03.01/actuals", json={"months": {"1": 10}}).status_code, 200)
        d = client.get("/api/budget").json()
        self.assertEqual(d["totals"]["supplement"], 30_000_000)
        with db.connect() as c:
            logged = c.execute("SELECT COUNT(*) FROM audit_log WHERE action='budget_decision'").fetchone()[0]
        self.assertEqual(logged, 2)

    def test_export_report_has_all_sheets_and_figures(self):
        client = self.client_for("ceo")
        client.post("/api/budget/plan/import", files={"file": ("p.xlsx", self.plan.read_bytes())})
        client.put("/api/budget/lines/9300.03.01/actuals", json={"months": {"1": 150_000_000}})
        client.post("/api/budget/lines/9300.03.01/decision", json={"kind": "supplement", "amount": 10_000_000, "note": "Giá xăng"})
        r = client.get("/api/budget/export?month=3")
        self.assertEqual(r.status_code, 200)
        self.assertIn("ngan-sach-2026-T3.xlsx", r.headers["content-disposition"])
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Tổng quan", "Chi tiết", "Cảnh báo", "Quyết định"])
        rows = list(wb["Chi tiết"].iter_rows(min_row=2, values_only=True))
        fuel = next(x for x in rows if x[0] == "9300.03.01")
        self.assertEqual((fuel[4], fuel[5], fuel[6], fuel[7], fuel[10]), (120_000_000, 10_000_000, 130_000_000, 150_000_000, "Vượt ngân sách"))
        self.assertEqual(wb["Cảnh báo"]["A2"].value, "9300.03.01")
        self.assertEqual(wb["Quyết định"]["C2"].value, "Duyệt bổ sung")

    def test_export_forbidden_for_viewer_and_needs_plan(self):
        self.assertEqual(self.client_for("viewer").get("/api/budget/export").status_code, 403)

    def test_export_without_plan_is_rejected(self):
        self.assertEqual(self.client_for("ceo").get("/api/budget/export").status_code, 400)

    def test_decision_forbidden_for_viewer(self):
        client = self.client_for("viewer")
        self.assertEqual(client.post("/api/budget/lines/x/decision", json={"kind": "block"}).status_code, 403)


if __name__ == "__main__":
    unittest.main()
