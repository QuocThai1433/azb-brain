"""Tiến độ mục tiêu doanh thu: % thực hiện, phân loại đạt / cần theo dõi / chậm / chưa có số, và chọn kỳ."""
import unittest
from datetime import date

from io import BytesIO

from openpyxl import load_workbook

from backend import analytics, export


def t(id_, period, target, actual, unit="Phòng A"):
    return {"id": id_, "period": period, "target_value": target, "actual_value": actual, "unit_name": unit}


class TargetProgressTest(unittest.TestCase):
    def test_status_thresholds(self):
        out = analytics.target_progress([
            t(1, "2026-10", 100, 100),   # đúng 100% -> đạt
            t(2, "2026-10", 100, 70),    # đúng 70% -> cần theo dõi
            t(3, "2026-10", 100, 69.9),  # dưới 70% -> chậm
            t(4, "2026-10", 100, None),  # chưa có số thực hiện
            t(5, "2026-10", None, 50),   # chưa có mục tiêu
            t(6, "2026-10", 0, 50),      # mục tiêu 0 không tính được %
        ])
        status = {x["id"]: x["status"] for x in out["items"]}
        self.assertEqual(status, {1: "ok", 2: "watch", 3: "behind", 4: "no_data", 5: "no_data", 6: "no_data"})
        self.assertEqual(out["summary"], {"behind": 1, "watch": 1, "ok": 1, "no_data": 3, "total": 6})

    def test_percent_is_rounded_to_one_decimal(self):
        out = analytics.target_progress([t(1, "2026-10", 300, 100)])
        self.assertEqual(out["items"][0]["pct"], 33.3)

    def test_defaults_to_latest_period_and_only_that_period(self):
        out = analytics.target_progress([t(1, "2026-09", 100, 100), t(2, "2026-10", 100, 50), t(3, "2026-Q4", 100, 80)])
        self.assertEqual(out["periods"], ["2026-Q4", "2026-10", "2026-09"])
        self.assertEqual(out["period"], "2026-Q4")
        self.assertEqual([x["id"] for x in out["items"]], [3])

    def test_explicit_period_and_unknown_period(self):
        rows = [t(1, "2026-09", 100, 100), t(2, "2026-10", 100, 50)]
        self.assertEqual([x["id"] for x in analytics.target_progress(rows, "2026-09")["items"]], [1])
        unknown = analytics.target_progress(rows, "2030")
        self.assertEqual(unknown["period"], "2026-10")  # kỳ không tồn tại -> lùi về kỳ mới nhất

    def test_sorted_with_attention_first(self):
        out = analytics.target_progress([
            t(1, "2026-10", 100, 100, "B"),
            t(2, "2026-10", 100, None, "C"),
            t(3, "2026-10", 100, 75, "D"),
            t(4, "2026-10", 100, 10, "E"),
        ])
        self.assertEqual([x["id"] for x in out["items"]], [4, 3, 1, 2])

    def test_empty(self):
        out = analytics.target_progress([])
        self.assertIsNone(out["period"])
        self.assertEqual(out["summary"]["total"], 0)

    def test_period_elapsed(self):
        e = analytics.period_elapsed
        self.assertAlmostEqual(e("2026", date(2026, 7, 2)), 183 / 365, places=3)
        self.assertEqual(e("2026", date(2026, 12, 31)), 1.0)
        self.assertEqual(e("2026-10", date(2026, 10, 16)), 16 / 31)
        self.assertEqual(e("2026-Q4", date(2026, 11, 15)), 46 / 92)
        self.assertEqual(e("2025", date(2026, 3, 1)), 1.0)   # kỳ đã qua
        self.assertEqual(e("2027", date(2026, 3, 1)), 0.0)   # kỳ chưa tới
        self.assertIsNone(e("FY26"))

    def test_forecast_and_expected(self):
        out = analytics.target_progress([t(1, "2026-10", 100, 50), t(2, "2026-10", 100, None)], today=date(2026, 10, 16))
        self.assertAlmostEqual(out["elapsed_pct"], 51.6, places=1)
        a, b = out["items"][0], out["items"][1]
        self.assertEqual(a["id"], 1)
        self.assertAlmostEqual(a["forecast_pct"], 96.9, places=1)
        self.assertEqual(a["forecast_value"], 97)
        self.assertIsNone(b["forecast_pct"])  # chưa có số thì không dự báo

    def test_no_forecast_at_very_start_of_period(self):
        out = analytics.target_progress([t(1, "2026", 100, 5)], today=date(2026, 1, 5))
        self.assertIsNone(out["items"][0]["forecast_pct"])

    def test_history_of_same_unit_and_metric(self):
        rows = [{**t(1, "2026-Q3", 100, 90), "unit_id": 5, "metric": "doanh_thu"},
                {**t(2, "2026-Q4", 100, 40), "unit_id": 5, "metric": "doanh_thu"},
                {**t(3, "2026-Q4", 100, 40), "unit_id": 6, "metric": "doanh_thu"}]
        item = next(x for x in analytics.target_progress(rows, "2026-Q4", date(2026, 11, 1))["items"] if x["id"] == 2)
        self.assertEqual([h["period"] for h in item["history"]], ["2026-Q3", "2026-Q4"])
        self.assertEqual(item["history"][0]["pct"], 90.0)

    def test_excel_export(self):
        rows = [{**t(1, "2026-10", 100, 50, "=HYPERLINK(\"x\")"), "unit_id": 1, "metric": "doanh_thu", "unit_label": "VNĐ"},
                {**t(2, "2026-10", 100, None, "B"), "unit_id": 2, "metric": "khac", "metric_name": "Khách mới"}]
        data = analytics.target_progress(rows, today=date(2026, 10, 16))
        ws = load_workbook(BytesIO(export.targets_xlsx(data))).active
        self.assertIn("2026-10", ws["A1"].value)
        self.assertEqual(ws["A5"].value, '=HYPERLINK("x")')   # văn bản, không phải công thức
        self.assertEqual(ws["A5"].data_type, "s")
        self.assertEqual((ws["D5"].value, ws["E5"].value, ws["F5"].value, ws["G5"].value), (100, 50, 50, 0.5))
        self.assertEqual(ws["J5"].value, "Chậm")
        self.assertEqual((ws["B6"].value, ws["J6"].value), ("Khách mới", "Chưa có số"))


if __name__ == "__main__":
    unittest.main()
