"""Ngân sách hoạt động năm 2026: kế hoạch (nhập từ file Excel "Bản duyệt ngân sách hoạt động"), số chi thực tế theo tháng
và cảnh báo vượt ngân sách.

Khác với budget.py (trần chi phí gọi AI): đây là ngân sách vận hành của công ty.
Mốc so sánh là cột "TỔNG DUYỆT" (S) của sheet TỔNG HỢP. Không dùng số năm 2025.

Nhập kế hoạch:  python -m backend.opbudget "đường dẫn file .xls/.xlsx"   (hoặc nút "Nhập kế hoạch" trên trang Ngân sách)
"""
import csv
import io
import math
import re
import sys
import unicodedata
from datetime import date

YEAR = 2026
WARN_PCT = 80        # >= 80% ngân sách đã dùng: sắp vượt
PACE_TOLERANCE = 1.1  # chi lũy kế vượt 110% mức theo tiến độ năm: báo "vượt tiến độ"

# Vị trí cột trong sheet TỔNG HỢP (A=0)
C_CODE, C_NAME, C_BASIC, C_EXTRA, C_TOTAL, C_INVEST = 0, 1, 11, 15, 18, 25
C_NOTES = (14, 16, 26, 27)  # ghi chú: cột O, Q, AA, AB (cột K là ghi chú của năm 2025 nên bỏ)

CODE_RE = re.compile(r"^\d{4}(\.\d{2})+$")


def _num(v) -> float:
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else 0.0
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return 0.0


def _title(s: str) -> str:
    s = re.sub(r"^TỔNG CỘNG\s+", "", re.sub(r"\s+", " ", str(s)).strip(), flags=re.I)
    return re.sub(r"\b(nnl|tgđ)\b", lambda m: m.group(1).upper(), s.capitalize()) if s.isupper() else s


def sheet_rows(path: str, prefer: str = "") -> list[list]:
    """Đọc một sheet thành danh sách dòng (.xls bằng xlrd, .xlsx bằng openpyxl); sheet có tên chứa `prefer` được ưu tiên."""
    p = str(path).lower()
    if p.endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(path)
        names = wb.sheet_names()
        sh = wb.sheet_by_name(next((n for n in names if prefer and prefer in n.upper()), names[0]))
        return [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    names = wb.sheetnames
    ws = wb[next((n for n in names if prefer and prefer in n.upper()), names[0])]
    rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close()  # Windows giữ khoá file nếu không đóng
    return rows


def parse_plan(path: str) -> list[dict]:
    rows = sheet_rows(path, "TỔNG HỢP")
    if not any(str(r[0]).strip().upper() == "MÃ CÔNG VIỆC" for r in rows[:30] if r):
        raise ValueError("Không phải file bản duyệt ngân sách (không thấy cột \"MÃ CÔNG VIỆC\" ở sheet TỔNG HỢP)")
    lines, section, group = [], None, ""
    for r in rows:
        r = list(r) + [""] * (28 - len(r))
        code, name = r[C_CODE], str(r[C_NAME]).strip()
        if isinstance(code, (int, float)):  # số thứ tự mục lớn: 1 = dòng tổng toàn công ty, 2..6 = nhóm chi phí
            section, group = (None, "") if code == 1 else (_title(name), "")
            continue
        code = str(code).strip()
        if CODE_RE.match(code):
            if section is None:
                continue
            notes = [re.sub(r"\s*\n\s*", "; ", r[c].strip()) for c in C_NOTES if isinstance(r[c], str)]
            lines.append({"code": code, "name": re.sub(r"\s+", " ", name), "section": section, "grp": group,
                          "plan_total": _num(r[C_TOTAL]), "basic": _num(r[C_BASIC]), "extra": _num(r[C_EXTRA]),
                          "invest": _num(r[C_INVEST]),
                          "note": " · ".join(n for n in notes if n and n != "Sheet 3") or None,
                          "sort_order": len(lines)})
        elif name and section and not name.upper().startswith("TỔNG CỘNG"):
            group = _title(name)
    if not lines:
        raise ValueError("Không đọc được hạng mục chi phí nào từ file")
    return lines


def import_plan(conn, path: str) -> dict:
    lines = parse_plan(path)
    codes = {l["code"] for l in lines}
    for l in lines:
        conn.execute("INSERT INTO budget_lines (code, name, section, grp, plan_total, basic, extra, invest, note, sort_order)"
                     " VALUES (:code,:name,:section,:grp,:plan_total,:basic,:extra,:invest,:note,:sort_order)"
                     " ON CONFLICT(code) DO UPDATE SET name=excluded.name, section=excluded.section, grp=excluded.grp,"
                     " plan_total=excluded.plan_total, basic=excluded.basic, extra=excluded.extra, invest=excluded.invest,"
                     " note=excluded.note, sort_order=excluded.sort_order", l)
    # hạng mục không còn trong file mới: chỉ xoá nếu chưa có số chi thực tế (không làm mất số liệu đã nhập)
    removed = 0
    for (code,) in conn.execute("SELECT code FROM budget_lines").fetchall():
        if code not in codes and not conn.execute("SELECT 1 FROM budget_actuals WHERE code=?", (code,)).fetchone():
            conn.execute("DELETE FROM budget_lines WHERE code=?", (code,))
            removed += 1
    return {"lines": len(lines), "removed": removed, "total": sum(l["plan_total"] for l in lines)}


# ---------- Số chi thực tế ----------
def is_blocked(conn, code: str) -> bool:
    r = conn.execute("SELECT kind FROM budget_decisions WHERE code=? AND kind IN ('block','unblock') ORDER BY id DESC LIMIT 1", (code,)).fetchone()
    return bool(r) and r[0] == "block"


def supplement_of(conn, code: str) -> float:
    return conn.execute("SELECT COALESCE(SUM(amount),0) FROM budget_decisions WHERE code=? AND kind='supplement'", (code,)).fetchone()[0]


def _year_spent(conn, code: str) -> float:
    return conn.execute("SELECT COALESCE(SUM(amount),0) FROM budget_actuals WHERE code=?", (code,)).fetchone()[0]


def save_actuals(conn, user_id: int, code: str, months: dict) -> int:
    """months: {tháng 1..12: số tiền}; ô trống/None = xoá số của tháng đó.
    Hạng mục đang bị chặn phát sinh: chỉ cho sửa giảm hoặc giữ nguyên tổng chi, không cho tăng."""
    if not conn.execute("SELECT 1 FROM budget_lines WHERE code=?", (code,)).fetchone():
        raise ValueError("Không có hạng mục này")
    before = _year_spent(conn, code)
    conn.execute("SAVEPOINT save_actuals")  # lỗi giữa chừng thì trả lại đúng hạng mục này, không ảnh hưởng hạng mục khác trong cùng lần nhập
    try:
        n = _write_actuals(conn, user_id, code, months)
        if is_blocked(conn, code) and _year_spent(conn, code) > before + 0.5:
            raise ValueError(f"Hạng mục {code} đang bị chặn phát sinh chi phí: Tổng Giám đốc cần duyệt bổ sung hoặc cho tiếp tục trước khi ghi thêm chi")
    except Exception:
        conn.execute("ROLLBACK TO save_actuals")
        conn.execute("RELEASE save_actuals")
        raise
    conn.execute("RELEASE save_actuals")
    return n


def _write_actuals(conn, user_id: int, code: str, months: dict) -> int:
    n = 0
    for k, v in months.items():
        m = int(k)
        if not 1 <= m <= 12:
            raise ValueError("Tháng phải từ 1 đến 12")
        if v is None or str(v).strip() == "":
            conn.execute("DELETE FROM budget_actuals WHERE code=? AND month=?", (code, m))
            continue
        amount = _num(v) if not isinstance(v, (int, float)) else float(v)
        if not math.isfinite(amount) or abs(amount) > 1e13:
            raise ValueError("Số tiền không hợp lệ")
        conn.execute("INSERT INTO budget_actuals (code, month, amount, updated_by, updated_at) VALUES (?,?,?,?,datetime('now','localtime'))"
                     " ON CONFLICT(code, month) DO UPDATE SET amount=excluded.amount, updated_by=excluded.updated_by, updated_at=excluded.updated_at",
                     (code, m, amount, user_id))
        n += 1
    return n


def _fold(s) -> str:
    s = unicodedata.normalize("NFD", str(s).lower().replace("đ", "d"))
    return re.sub(r"[^a-z0-9]+", " ", "".join(c for c in s if unicodedata.category(c) != "Mn")).strip()


def _month_of(header) -> int | None:
    if isinstance(header, float) and header == int(header) and 1 <= header <= 12:
        return int(header)
    m = re.fullmatch(r"(?:t|thang)? ?(\d{1,2})", _fold(header))
    return int(m.group(1)) if m and 1 <= int(m.group(1)) <= 12 else None


def read_table(data: bytes, filename: str) -> list[list]:
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        delim = ";" if text.split("\n", 1)[0].count(";") > text.split("\n", 1)[0].count(",") else ","
        return [r for r in csv.reader(io.StringIO(text), delimiter=delim)]
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows = [["" if c is None else c for c in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    return rows


def import_actuals(conn, user_id: int, data: bytes, filename: str) -> dict:
    """Hai dạng file: (1) mỗi dòng một hạng mục, các cột T1..T12; (2) mỗi dòng một khoản: Mã, Tháng, Số tiền.
    Ô trống bị bỏ qua (không xoá số đã có)."""
    table = read_table(data, filename)
    head_i = next((i for i, r in enumerate(table[:10]) if any(_fold(c) in ("ma", "ma cong viec", "ma hang muc") for c in r)), None)
    if head_i is None:
        raise ValueError("Không thấy cột \"Mã\" ở 10 dòng đầu của file")
    head = table[head_i]
    ci = next(i for i, c in enumerate(head) if _fold(c) in ("ma", "ma cong viec", "ma hang muc"))
    month_cols = {i: _month_of(c) for i, c in enumerate(head) if _month_of(c)}
    mi = next((i for i, c in enumerate(head) if _fold(c) in ("thang", "ky")), None)
    ai = next((i for i, c in enumerate(head) if _fold(c) in ("so tien", "thuc chi", "chi thuc te", "amount")), None)
    if not month_cols and (mi is None or ai is None):
        raise ValueError("File cần các cột T1..T12, hoặc hai cột \"Tháng\" và \"Số tiền\"")
    known = {c for (c,) in conn.execute("SELECT code FROM budget_lines")}
    per_code: dict[str, dict] = {}
    unknown: set[str] = set()
    for r in table[head_i + 1:]:
        if ci >= len(r) or not str(r[ci]).strip():
            continue
        code = str(r[ci]).strip()
        if code not in known:
            unknown.add(code)
            continue
        cells = {month_cols[i]: r[i] for i in month_cols if i < len(r)} if month_cols else \
            {_month_of(r[mi]): r[ai]} if mi < len(r) and ai < len(r) and _month_of(r[mi]) else {}
        for m, v in cells.items():
            if str(v).strip() != "":
                per_code.setdefault(code, {})[m] = v
    saved, blocked, done = 0, [], 0
    for c, months in per_code.items():
        try:
            saved += save_actuals(conn, user_id, c, months)
            done += 1
        except ValueError:  # hạng mục đang bị chặn phát sinh: bỏ qua riêng hạng mục đó
            blocked.append(c)
    return {"lines": done, "entries": saved, "unknown": sorted(unknown)[:20], "blocked": blocked}


def decide(conn, user_id: int, code: str, kind: str, amount=None, note: str = "") -> dict:
    """Quyết định của Ban TGĐ: supplement (duyệt thêm ngân sách, tự gỡ chặn nếu đang chặn), block, unblock."""
    if not conn.execute("SELECT 1 FROM budget_lines WHERE code=?", (code,)).fetchone():
        raise ValueError("Không có hạng mục này")
    note = (note or "").strip()[:500]
    blocked = is_blocked(conn, code)
    if kind == "supplement":
        amount = _num(amount) if not isinstance(amount, (int, float)) else float(amount)
        if not math.isfinite(amount) or not 0 < amount <= 1e13:
            raise ValueError("Số tiền bổ sung phải lớn hơn 0")
        if not note:
            raise ValueError("Cần ghi lý do duyệt bổ sung")
    elif kind == "block":
        if blocked:
            raise ValueError("Hạng mục đã bị chặn phát sinh")
        amount = None
    elif kind == "unblock":
        if not blocked:
            raise ValueError("Hạng mục không bị chặn")
        amount = None
    else:
        raise ValueError("Quyết định không hợp lệ")
    conn.execute("INSERT INTO budget_decisions (code, kind, amount, note, decided_by) VALUES (?,?,?,?,?)", (code, kind, amount, note or None, user_id))
    if kind == "supplement" and blocked:
        conn.execute("INSERT INTO budget_decisions (code, kind, note, decided_by) VALUES (?,?,?,?)",
                     (code, "unblock", "Tự gỡ chặn khi duyệt bổ sung ngân sách", user_id))
    return {"blocked": is_blocked(conn, code), "supplement": supplement_of(conn, code)}


def template_xlsx(conn) -> bytes:
    """File mẫu nhập số chi: mỗi hạng mục một dòng, đã điền sẵn số đã nhập (nếu có) để sửa rồi nhập lại."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chi thực tế 2026"
    ws.append(["Mã", "Hạng mục", "Ngân sách 2026"] + [f"T{m}" for m in range(1, 13)])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="004C97")
    acts: dict[str, dict] = {}
    for code, m, amt in conn.execute("SELECT code, month, amount FROM budget_actuals"):
        acts.setdefault(code, {})[m] = amt
    for code, name, plan in conn.execute("SELECT code, name, plan_total FROM budget_lines ORDER BY sort_order"):
        ws.append([code, name, plan] + [acts.get(code, {}).get(m) for m in range(1, 13)])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 12, 60
    for col in "CDEFGHIJKLMNO":
        ws.column_dimensions[col].width = 16
        for cell in ws[col][1:]:
            cell.number_format = "#,##0"
    ws.freeze_panes = "D2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- Dashboard ----------
def line_status(plan: float, spent: float, expected: float) -> str:
    """over: vượt ngân sách | near: >= 80% | pace: chi nhanh hơn tiến độ năm | ok | none: chưa có ngân sách và chưa chi."""
    if spent > plan + 0.5:
        return "over"
    if plan <= 0:
        return "none"
    if spent >= plan * WARN_PCT / 100:
        return "near"
    if spent > expected * PACE_TOLERANCE and spent > 0:
        return "pace"
    return "ok"


def _agg(items: list[dict], as_of: int) -> dict:
    plan, spent = sum(i["plan"] for i in items), sum(i["spent"] for i in items)
    return {"plan": plan, "spent": spent, "remaining": plan - spent,
            "percent": round(spent / plan * 100, 1) if plan > 0 else (100.0 if spent > 0 else 0.0),
            "status": line_status(plan, spent, plan * as_of / 12),
            "over": sum(1 for i in items if i["status"] == "over"), "near": sum(1 for i in items if i["status"] == "near"),
            "pace": sum(1 for i in items if i["status"] == "pace"), "blocked": sum(1 for i in items if i["blocked"]),
            "supplement": sum(i["supplement"] for i in items)}


def dashboard(conn, month: int | None = None, today: date | None = None) -> dict:
    lines = [dict(r) for r in conn.execute("SELECT * FROM budget_lines ORDER BY sort_order")]
    if not lines:
        return {"year": YEAR, "has_plan": False}
    acts: dict[str, list] = {}
    last_with_data = 0
    for code, m, amt in conn.execute("SELECT code, month, amount FROM budget_actuals"):
        acts.setdefault(code, [None] * 12)[m - 1] = amt
        last_with_data = max(last_with_data, m)
    supp: dict[str, float] = {}
    for code, amt in conn.execute("SELECT code, SUM(amount) FROM budget_decisions WHERE kind='supplement' GROUP BY code"):
        supp[code] = amt or 0.0
    blocked = {c for (c,) in conn.execute(
        "SELECT code FROM budget_decisions d WHERE kind IN ('block','unblock') AND id = "
        "(SELECT MAX(id) FROM budget_decisions WHERE code = d.code AND kind IN ('block','unblock')) AND kind = 'block'")}
    today = today or date.today()
    as_of = month if month and 1 <= month <= 12 else (last_with_data or (today.month if today.year == YEAR else 12 if today.year > YEAR else 1))
    sections: dict[str, dict] = {}
    for l in lines:
        months = acts.get(l["code"], [None] * 12)
        spent = sum(v or 0 for v in months[:as_of])
        plan = l["plan_total"] + supp.get(l["code"], 0.0)
        item = {"code": l["code"], "name": l["name"], "plan": plan, "base_plan": l["plan_total"], "supplement": supp.get(l["code"], 0.0),
                "blocked": l["code"] in blocked, "spent": spent, "remaining": plan - spent,
                "percent": round(spent / plan * 100, 1) if plan > 0 else (100.0 if spent > 0 else 0.0),
                "expected": plan * as_of / 12, "over_by": max(0.0, spent - plan), "note": l["note"], "months": months,
                "section": l["section"], "grp": l["grp"], "year_spent": sum(v or 0 for v in months)}
        item["status"] = line_status(plan, spent, item["expected"])
        sections.setdefault(l["section"], {}).setdefault(l["grp"], []).append(item)
    all_items = [i for g in sections.values() for items in g.values() for i in items]
    out_sections = []
    for name, groups in sections.items():
        sec_items = [i for items in groups.values() for i in items]
        out_sections.append({"name": name, **_agg(sec_items, as_of),
                             "groups": [{"name": g, **_agg(items, as_of), "lines": items} for g, items in groups.items()]})
    order = {"over": 0, "near": 1, "pace": 3}
    # hạng mục đang bị chặn luôn nằm trong danh sách để TGĐ còn thấy mà cho tiếp tục
    alerts = sorted((i for i in all_items if i["status"] in order or i["blocked"]),
                    key=lambda i: (order.get(i["status"], 2), -i["over_by"], -i["percent"]))
    names = {i["code"]: i["name"] for i in all_items}
    decisions = [dict(r) | {"name": names.get(r["code"], "")} for r in conn.execute(
        "SELECT d.id, d.code, d.kind, d.amount, d.note, d.decided_at, COALESCE(u.full_name, u.username) AS by_name "
        "FROM budget_decisions d LEFT JOIN users u ON u.id = d.decided_by ORDER BY d.id DESC LIMIT 15")]
    monthly, cum_spent = [], 0.0
    plan_total = sum(i["plan"] for i in all_items)
    for m in range(1, 13):
        s = sum((acts.get(i["code"], [None] * 12)[m - 1] or 0) for i in all_items)
        cum_spent += s
        monthly.append({"month": m, "spent": s, "cum_spent": cum_spent, "cum_plan": plan_total * m / 12, "has_data": m <= last_with_data})
    return {"year": YEAR, "has_plan": True, "as_of": as_of, "last_data_month": last_with_data,
            "warn_pct": WARN_PCT, "totals": _agg(all_items, as_of) | {"expected": plan_total * as_of / 12},
            "sections": out_sections, "alerts": alerts, "monthly": monthly, "decisions": decisions}


STATUS_LABEL = {"over": "Vượt ngân sách", "near": "Sắp vượt", "pace": "Vượt tiến độ", "ok": "Trong mức", "none": "Chưa có ngân sách"}
KIND_LABEL = {"supplement": "Duyệt bổ sung", "block": "Chặn phát sinh", "unblock": "Cho tiếp tục"}


def export_xlsx(conn, month: int | None = None) -> bytes:
    """Báo cáo ngân sách để trình TGĐ: tổng quan theo nhóm, chi tiết từng hạng mục, cảnh báo và quyết định."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    d = dashboard(conn, month)
    if not d["has_plan"]:
        raise ValueError("Chưa có kế hoạch ngân sách để xuất")
    m, T = d["as_of"], d["totals"]
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="004C97")
    fills = {"over": "FDECEC", "near": "FFF1E6", "pace": "FFF8E1"}
    money = "#,##0"

    def sheet(ws, header, rows, widths, money_cols=(), pct_cols=(), status_col=None):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.column in money_cols:
                    c.number_format = money
                elif c.column in pct_cols:
                    c.number_format = '0.0"%"'
            if status_col:
                st = {v: k for k, v in STATUS_LABEL.items()}.get(row[status_col - 1].value)
                if st in fills:
                    for c in row:
                        c.fill = PatternFill("solid", fgColor=fills[st])
        ws.freeze_panes = "A2"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tổng quan"
    ws.append([f"NGÂN SÁCH HOẠT ĐỘNG {d['year']} - lũy kế đến tháng {m}"])
    ws["A1"].font = Font(bold=True, size=14, color="004C97")
    ws.append([f"Ngày xuất: {date.today():%d/%m/%Y}. Mốc so sánh: ngân sách tổng duyệt + phần Tổng Giám đốc duyệt bổ sung."])
    ws.append([])
    for label, val in (("Ngân sách hiệu lực", T["plan"]), ("Trong đó đã duyệt bổ sung", T["supplement"]), ("Đã chi lũy kế", T["spent"]),
                       ("Còn lại", T["remaining"]), ("Mức theo tiến độ năm", T["expected"])):
        ws.append([label, val])
        ws.cell(ws.max_row, 2).number_format = money
    ws.append(["Đã sử dụng (%)", T["percent"]])
    ws.cell(ws.max_row, 2).number_format = '0.0"%"'
    ws.append(["Hạng mục vượt / sắp vượt / vượt tiến độ / đang bị chặn", f"{T['over']} / {T['near']} / {T['pace']} / {T['blocked']}"])
    ws.append([])
    base = ws.max_row + 1
    ws.append(["Nhóm chi phí", "Ngân sách", "Đã chi", "Còn lại", "% đã dùng", "Trạng thái"])
    for c in ws[base]:
        c.font, c.fill = head_font, head_fill
    for sec in d["sections"]:
        ws.append([sec["name"], sec["plan"], sec["spent"], sec["remaining"], sec["percent"], STATUS_LABEL[sec["status"]]])
        for col in (2, 3, 4):
            ws.cell(ws.max_row, col).number_format = money
        ws.cell(ws.max_row, 5).number_format = '0.0"%"'
    ws.column_dimensions["A"].width, ws.column_dimensions["F"].width = 58, 20
    for col in "BCDE":
        ws.column_dimensions[col].width = 18

    lines = [i for sec in d["sections"] for g in sec["groups"] for i in g["lines"]]
    sheet(wb.create_sheet("Chi tiết"),
          ["Mã", "Hạng mục", "Nhóm chi phí", "Nhóm con", "Ngân sách gốc", "Bổ sung", "Ngân sách hiệu lực", f"Đã chi đến T{m}", "Còn lại", "% đã dùng", "Trạng thái", "Chặn phát sinh", "Ghi chú"],
          [[i["code"], i["name"], i["section"], i["grp"], i["base_plan"], i["supplement"], i["plan"], i["spent"], i["remaining"], i["percent"],
            STATUS_LABEL[i["status"]], "Có" if i["blocked"] else "", i["note"] or ""] for i in lines],
          [12, 56, 34, 28, 16, 14, 18, 18, 16, 11, 18, 13, 44], money_cols=(5, 6, 7, 8, 9), pct_cols=(10,), status_col=11)
    sheet(wb.create_sheet("Cảnh báo"),
          ["Mã", "Hạng mục", "Trạng thái", "Ngân sách", "Đã chi", "Vượt", "% đã dùng", "Chặn phát sinh"],
          [[a["code"], a["name"], STATUS_LABEL[a["status"]], a["plan"], a["spent"], a["over_by"], a["percent"], "Có" if a["blocked"] else ""] for a in d["alerts"]],
          [12, 56, 18, 18, 18, 16, 11, 14], money_cols=(4, 5, 6), pct_cols=(7,), status_col=3)
    sheet(wb.create_sheet("Quyết định"),
          ["Thời điểm", "Hạng mục", "Quyết định", "Số tiền", "Lý do", "Người quyết định"],
          [[x["decided_at"], f"{x['code']} {x['name']}", KIND_LABEL[x["kind"]], x["amount"], x["note"] or "", x["by_name"] or ""] for x in d["decisions"]],
          [18, 56, 18, 16, 50, 26], money_cols=(4,))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    from .db import connect, init_db
    if len(sys.argv) != 2:
        sys.exit('Cách dùng: python -m backend.opbudget "đường dẫn file bản duyệt ngân sách .xls/.xlsx"')
    init_db()
    with connect() as conn:
        res = import_plan(conn, sys.argv[1])
    print(f"Đã nhập {res['lines']} hạng mục, tổng ngân sách duyệt {res['total']:,.0f} đ; bỏ {res['removed']} hạng mục cũ.")
